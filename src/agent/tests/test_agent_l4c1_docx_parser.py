"""Layer 4C1 docx parser tests."""

from __future__ import annotations

from io import BytesIO
import unittest

from agent.documents.models import DocumentSource
from agent.documents.parsers.docx import PythonDocxParser, docx_parser_available


@unittest.skipUnless(docx_parser_available(), "python-docx optional dependency unavailable")
class AgentLayer4C1DocxParserTests(unittest.TestCase):
    def setUp(self) -> None:
        self.parser = PythonDocxParser()
        self.source = DocumentSource(
            doc_id="docx123456789abc",
            doc_name="sample.docx",
            doc_type="docx",
            byte_size=0,
            ingested_at=1,
        )

    def _docx_bytes(self) -> bytes:
        from docx import Document  # type: ignore[import-not-found]

        document = Document()
        document.add_heading("Article 3.2", level=1)
        document.add_paragraph("If account is dormant then flag review.")
        document.add_paragraph("Alice owns account.")
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    def test_parse_docx_extracts_heading_as_section_label(self) -> None:
        segments = self.parser.parse(self.source, self._docx_bytes())
        self.assertGreaterEqual(len(segments), 3)
        self.assertEqual(segments[0].section_label, "Article 3.2")
        self.assertEqual(segments[1].section_label, "Article 3.2")
        self.assertEqual(segments[1].pattern_type, "if_then")


if __name__ == "__main__":
    unittest.main()
