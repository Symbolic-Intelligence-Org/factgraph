"""Layer 4C1 document staging integration tests."""

from __future__ import annotations

from io import BytesIO
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from factpy.agent import (
    AgentCheckpointStore,
    AgentScope,
    AgentSession,
    CandidatePayloadCache,
    DocumentStaging,
    DraftManager,
    EvaluateTools,
    ReadReviewOrchestrator,
    RuntimeBootstrapSpec,
    build_layer3a_tool_registry,
)
from factpy.agent.documents.parsers.docx import docx_parser_available
from factpy.agent.documents.parsers.pdf import pdf_parser_available
from factpy.agent.tools._runtime_api import LocalRuntimeAPI
from factpy.agent.tools.explain import ExplainTools
from factpy.agent.tools.kg_read import KGReadTools
from factpy.service.runtime_v1 import (
    close_runtime_session,
    get_runtime_session,
    open_runtime_session,
    reset_runtime_sessions_for_tests,
)
from factpy.tests._test_helpers import _schema_ir


def _open_session(open_dto: dict[str, object]) -> str:
    resp = open_runtime_session(open_dto)
    assert resp["ok"], resp
    return resp["session"]["session_id"]


class AgentLayer4C1StagingTests(unittest.TestCase):
    def setUp(self) -> None:
        reset_runtime_sessions_for_tests()
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db_path = Path(self.tmp.name) / "agent_l4c1.sqlite3"
        self.ledger_path = str(Path(self.tmp.name) / "runtime-ledger.db")
        self.open_dto = {"schema_ir": _schema_ir(), "ledger_path": self.ledger_path}
        self.runtime_session_id = _open_session(self.open_dto)

        self.agent_session = AgentSession(scope=AgentScope(agent_id="agent-l4c1"))
        self.agent_session.bind_runtime_session(
            self.runtime_session_id,
            bootstrap_spec=RuntimeBootstrapSpec.from_open_dto(self.open_dto),
            burr_db_path=str(self.db_path),
        )
        self.draft_manager = DraftManager()
        self.cache = CandidatePayloadCache(self.db_path)
        self.checkpoint = AgentCheckpointStore(self.db_path)
        self.runtime_api = LocalRuntimeAPI()
        self.kg = KGReadTools(runtime_api=self.runtime_api)
        self.explain = ExplainTools(runtime_api=self.runtime_api)
        self.evaluate_tools = EvaluateTools(
            runtime_api=self.runtime_api,
            candidate_cache=self.cache,
            explain_tools=self.explain,
            session=self.agent_session,
        )
        self.document_staging = DocumentStaging()
        self.orchestrator = ReadReviewOrchestrator(
            session=self.agent_session,
            draft_manager=self.draft_manager,
            kg_read_tools=self.kg,
            explain_tools=self.explain,
            evaluate_tools=self.evaluate_tools,
            candidate_cache=self.cache,
            checkpoint_store=self.checkpoint,
            document_staging=self.document_staging,
        )

    def tearDown(self) -> None:
        self.cache.close()
        self.checkpoint.close()
        if get_runtime_session(self.runtime_session_id).get("ok") is True:
            close_runtime_session(self.runtime_session_id)
        reset_runtime_sessions_for_tests()

    def test_doc_id_is_idempotent_and_segments_are_deterministic(self) -> None:
        payload = b"# Intro\nIf user is vip then notify.\n\nNarrative."
        first = self.document_staging.stage_document(doc_name="sample.md", content=payload)
        second = self.document_staging.stage_document(doc_name="sample.md", content=payload)
        self.assertEqual(first.source.doc_id, second.source.doc_id)
        self.assertEqual(first.segments, second.segments)

    def test_supported_formats_follow_dependency_registration(self) -> None:
        formats = self.document_staging.list_supported_formats()
        self.assertIn("txt", formats)
        self.assertIn("md", formats)
        if docx_parser_available():
            self.assertIn("docx", formats)
        if pdf_parser_available():
            self.assertIn("pdf", formats)
        else:
            self.assertNotIn("pdf", formats)

    def test_missing_pdf_dependency_returns_unsupported_format(self) -> None:
        if pdf_parser_available():
            self.skipTest("pdf optional dependencies available; unsupported path not applicable")
        result = self.document_staging.stage_document(doc_name="sample.pdf", content=b"%PDF-1.4")
        self.assertEqual(result.error_kind, "unsupported_format")

    @unittest.skipUnless(docx_parser_available(), "python-docx optional dependency unavailable")
    def test_stage_docx_returns_segments(self) -> None:
        from docx import Document  # type: ignore[import-not-found]

        document = Document()
        document.add_heading("Section 4.1", level=1)
        document.add_paragraph("If account is dormant then flag review.")
        stream = BytesIO()
        document.save(stream)

        result = self.document_staging.stage_document(
            doc_name="sample.docx",
            content=stream.getvalue(),
        )
        self.assertEqual(result.source.doc_type, "docx")
        self.assertGreaterEqual(len(result.segments), 2)

    def test_orchestrator_stage_document_does_not_touch_session(self) -> None:
        before = self.agent_session.last_active_at
        result = self.orchestrator.stage_document(
            doc_name="sample.txt",
            content=b"Article 3.2\nIf user is vip then notify.\n\nNarrative.",
        )
        self.assertEqual(result.source.doc_type, "txt")
        self.assertEqual(self.agent_session.last_active_at, before)

    def test_build_layer3a_tool_registry_exposes_forty_five_tools(self) -> None:
        tool_registry = build_layer3a_tool_registry(orchestrator=self.orchestrator)
        self.assertEqual(len(tool_registry), 45)
        self.assertIn("stage_document", tool_registry)
        self.assertIn("list_supported_document_formats", tool_registry)
        self.assertIn("create_document_bundle", tool_registry)
        self.assertIn("commit_bundle", tool_registry)
        self.assertIn("extract_from_segment", tool_registry)
        self.assertIn("extract_and_create_bundle", tool_registry)


if __name__ == "__main__":
    unittest.main()
